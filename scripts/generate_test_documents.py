
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "documents"
GROUND_TRUTH = ROOT / "results" / "ground_truth.json"

invoice = {
    "vendor": "CraftWorks Supplies",
    "invoice_number": "CW-2026-1042",
    "invoice_date": "2026-10-15",
    "customer": "ZeonAI Labs",
    "items": [
        {"description": "Notebook", "quantity": 10, "unit_price": 50.00},
        {"description": "Pen Set", "quantity": 5, "unit_price": 80.00},
    ],
    "subtotal": 900.00,
    "tax_rate_percent": 18,
    "tax": 162.00,
    "total": 1062.00,
}

lines = [
    "INVOICE",
    f"Vendor: {invoice['vendor']}",
    f"Invoice Number: {invoice['invoice_number']}",
    f"Invoice Date: {invoice['invoice_date']}",
    f"Customer: {invoice['customer']}",
    "",
    "Item                 Qty    Unit Price    Amount",
    "Notebook              10       50.00      500.00",
    "Pen Set                5       80.00      400.00",
    "",
    "Subtotal: 900.00",
    "Tax (18%): 162.00",
    "Total: 1062.00",
]


def ensure_dirs():
    for folder in [
        "digital_pdf",
        "scanned_pdf",
        "docx",
        "xlsx",
        "pptx",
        "images",
    ]:
        (DOCS / folder).mkdir(parents=True, exist_ok=True)


def create_digital_pdf():
    path = DOCS / "digital_pdf" / "synthetic_invoice.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    y = 790

    for line in lines:
        pdf.drawString(55, y, line)
        y -= 25

    pdf.save()


def create_scanned_pdf():
    image_path = DOCS / "invoices" / "images" / "invoice_scan.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)

    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)

    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 28)
        heading_font = ImageFont.truetype("DejaVuSans.ttf", 40)
    except OSError:
        font = ImageFont.load_default()
        heading_font = font

    y = 100
    for index, line in enumerate(lines):
        draw.text(
            (90, y),
            line,
            fill="black",
            font=heading_font if index == 0 else font,
        )
        y += 65

    image.save(image_path)

    pdf_path = DOCS / "scanned_pdf" / "synthetic_invoice_scan.pdf"
    image.save(pdf_path, "PDF", resolution=150)


def create_docx():
    doc = Document()
    doc.add_heading("INVOICE", level=1)
    doc.add_paragraph(f"Vendor: {invoice['vendor']}")
    doc.add_paragraph(f"Invoice Number: {invoice['invoice_number']}")
    doc.add_paragraph(f"Invoice Date: {invoice['invoice_date']}")
    doc.add_paragraph(f"Customer: {invoice['customer']}")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = ["Item", "Quantity", "Unit Price", "Amount"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value

    for item in invoice["items"]:
        row = table.add_row().cells
        row[0].text = item["description"]
        row[1].text = str(item["quantity"])
        row[2].text = f"{item['unit_price']:.2f}"
        row[3].text = (
            f"{item['quantity'] * item['unit_price']:.2f}"
        )

    doc.add_paragraph(f"Subtotal: {invoice['subtotal']:.2f}")
    doc.add_paragraph(f"Tax (18%): {invoice['tax']:.2f}")
    doc.add_paragraph(f"Total: {invoice['total']:.2f}")

    doc.save(DOCS / "docx" / "synthetic_invoice.docx")


def create_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"

    rows = [
        ["INVOICE"],
        ["Vendor", invoice["vendor"]],
        ["Invoice Number", invoice["invoice_number"]],
        ["Invoice Date", invoice["invoice_date"]],
        ["Customer", invoice["customer"]],
        [],
        ["Item", "Quantity", "Unit Price", "Amount"],
    ]

    for item in invoice["items"]:
        rows.append([
            item["description"],
            item["quantity"],
            item["unit_price"],
            item["quantity"] * item["unit_price"],
        ])

    rows.extend([
        [],
        ["Subtotal", invoice["subtotal"]],
        ["Tax (18%)", invoice["tax"]],
        ["Total", invoice["total"]],
    ])

    for row in rows:
        ws.append(row)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

    # Second sheet tests multi-sheet workbook handling.
    notes = wb.create_sheet("Notes")
    notes.append(["Field", "Expected Value"])
    notes.append(["Currency", "INR"])
    notes.append(["Invoice Status", "Unpaid"])

    wb.save(DOCS / "xlsx" / "synthetic_invoice.xlsx")


def create_pptx():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "INVOICE"

    box = slide.shapes.add_textbox(
        Inches(0.7), Inches(1.3), Inches(8.5), Inches(1.6)
    )
    tf = box.text_frame
    tf.text = f"Vendor: {invoice['vendor']}"
    tf.add_paragraph().text = (
        f"Invoice Number: {invoice['invoice_number']}"
    )
    tf.add_paragraph().text = f"Date: {invoice['invoice_date']}"
    tf.add_paragraph().text = f"Customer: {invoice['customer']}"

    table_shape = slide.shapes.add_table(
        3, 4, Inches(0.7), Inches(3.2), Inches(8.5), Inches(1.2)
    )
    table = table_shape.table

    data = [
        ["Item", "Quantity", "Unit Price", "Amount"],
        ["Notebook", "10", "50.00", "500.00"],
        ["Pen Set", "5", "80.00", "400.00"],
    ]

    for r, row in enumerate(data):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    total_box = slide.shapes.add_textbox(
        Inches(0.7), Inches(5.0), Inches(6), Inches(1)
    )
    total_box.text_frame.text = (
        f"Subtotal: 900.00 | Tax: 162.00 | Total: 1062.00"
    )

    prs.save(DOCS / "pptx" / "synthetic_invoice.pptx")


def create_png():
    path = DOCS / "images" / "synthetic_invoice.png"
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)

    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 30)
        heading = ImageFont.truetype("DejaVuSans.ttf", 44)
    except OSError:
        font = ImageFont.load_default()
        heading = font

    y = 100
    for index, line in enumerate(lines):
        draw.text(
            (80, y),
            line,
            fill="black",
            font=heading if index == 0 else font,
        )
        y += 75

    image.save(path)


def main():
    ensure_dirs()
    create_digital_pdf()
    create_scanned_pdf()
    create_docx()
    create_xlsx()
    create_pptx()
    create_png()

    GROUND_TRUTH.parent.mkdir(parents=True, exist_ok=True)
    GROUND_TRUTH.write_text(
        json.dumps(invoice, indent=2),
        encoding="utf-8",
    )

    print("Synthetic test documents created.")
    print(f"Ground truth: {GROUND_TRUTH}")
    print("Created: digital PDF, scanned PDF, DOCX, XLSX, PPTX, PNG")


if __name__ == "__main__":
    main()
